# ─── RadioScript · ui/ventana_reportes.py ────────────────────────────────────
import tkinter as tk
from tkinter import font as tkfont, messagebox, filedialog, ttk
import os
import re
import subprocess
import sys
import shutil
import time
import threading
from datetime import datetime
from db.conexion import Conexion
from db.repositorios.transcripciones import InformesRepo
from config import (
    BG_DARK, BG_CARD, BG_ITEM, BG_HOVER,
    ACCENT_BLUE, ACCENT_GREEN, ACCENT_RED, ACCENT_YELLOW,
    TEXT_PRIMARY, TEXT_MUTED, TEXT_ACCENT,
    BORDER, ICONO_OK
)


def _tiene_ffplay() -> bool:
    return shutil.which("ffplay") is not None


def _duracion_audio(path: str) -> float:
    """Obtiene la duración en segundos usando ffprobe."""
    try:
        resultado = subprocess.run(
            ["ffprobe", "-v", "error", "-show_entries", "format=duration",
             "-of", "default=noprint_wrappers=1:nokey=1", path],
            capture_output=True, text=True, timeout=10
        )
        return float(resultado.stdout.strip())
    except Exception:
        return 0.0


def _seg_a_tiempo(seg: float) -> str:
    seg = int(seg)
    m, s = divmod(seg, 60)
    return f"{m}:{s:02d}"


def _guardar_docx(informe_texto: str, nombre_archivo: str, carpeta: str) -> str:
    try:
        from docx import Document
        from docx.shared import Pt, Inches
        doc = Document()
        for section in doc.sections:
            section.top_margin    = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin   = Inches(1.25)
            section.right_margin  = Inches(1.25)
        style = doc.styles["Normal"]
        style.font.name = "Arial"
        style.font.size = Pt(11)
        for linea in informe_texto.split("\n"):
            linea = linea.rstrip()
            p   = doc.add_paragraph()
            run = p.add_run(linea)
            run.bold = linea.strip().endswith(":") or linea.startswith("*")
            run.font.size = Pt(11)
            p.paragraph_format.space_after = Pt(2)
        ruta = os.path.join(carpeta, nombre_archivo + ".docx")
        doc.save(ruta)
        return ruta
    except ImportError:
        ruta = os.path.join(carpeta, nombre_archivo + ".txt")
        with open(ruta, "w", encoding="utf-8") as f:
            f.write(informe_texto)
        return ruta


def _nombre_archivo(resultado: dict) -> str:
    partes = []
    if resultado.get("tipo_estudio"):
        partes.append(resultado["tipo_estudio"].replace(" ", "_"))
    if resultado.get("nombre_paciente"):
        nombre = re.sub(r'[^\w\s-]', '', resultado["nombre_paciente"])
        partes.append(nombre.replace(" ", "_"))
    partes.append(datetime.now().strftime("%Y%m%d"))
    nombre = "_".join(p for p in partes if p)
    return re.sub(r'[<>:"/\\|?*]', '_', nombre) or "reporte_sin_nombre"


def _textos_diferentes(a: str, b: str) -> bool:
    limpiar = lambda t: re.sub(r'\s+', ' ', t.strip().upper())
    return limpiar(a) != limpiar(b)


class VentanaReportes(tk.Toplevel):
    def __init__(self, parent, reportes: list):
        super().__init__(parent)
        self.title("📄  Reportes Generados — RadioScript")
        self.configure(bg=BG_DARK)
        self.resizable(True, True)
        self.minsize(860, 660)

        self._reportes = reportes
        self._indice   = 0
        self._fuentes  = self._crear_fuentes()
        self._carpeta  = self._carpeta_documentos()

        # Audio
        self._ffplay_proc    = None
        self._audio_playing  = False
        self._tiene_ffplay   = _tiene_ffplay()
        self._audio_inicio   = 0.0   # tiempo real cuando empezó a reproducir
        self._audio_duracion = 0.0   # duración total del audio actual
        self._seek_activo    = False  # el usuario está arrastrando la barra

        self._construir_ui()
        self._cargar_reporte(0)

        self.update_idletasks()
        w, h = 1020, 760
        px, py = parent.winfo_rootx(), parent.winfo_rooty()
        pw, ph = parent.winfo_width(), parent.winfo_height()
        self.geometry(f"{w}x{h}+{px + pw//2 - w//2}+{py + ph//2 - h//2}")

        self.transient(parent)
        self.deiconify()
        self.after(100, self.grab_set)
        self.protocol("WM_DELETE_WINDOW", self._cerrar)

    def _crear_fuentes(self):
        return {
            "titulo": tkfont.Font(family="DejaVu Sans", size=13, weight="bold"),
            "label":  tkfont.Font(family="DejaVu Sans", size=10, weight="bold"),
            "body":   tkfont.Font(family="DejaVu Sans", size=10),
            "editor": tkfont.Font(family="DejaVu Sans Mono", size=11),
            "small":  tkfont.Font(family="DejaVu Sans", size=9),
            "mono":   tkfont.Font(family="DejaVu Sans Mono", size=8),
        }

    def _carpeta_documentos(self) -> str:
        base = os.path.join(os.path.expanduser("~"), "Documentos", "RadioScript")
        os.makedirs(base, exist_ok=True)
        return base

    # ─── UI ───────────────────────────────────────────────────────────────────
    def _construir_ui(self):
        # Header
        header = tk.Frame(self, bg=BG_CARD, pady=10)
        header.pack(fill=tk.X)
        tk.Label(header, text="📄  Editor de Reportes",
                 font=self._fuentes["titulo"],
                 bg=BG_CARD, fg=TEXT_PRIMARY).pack(side=tk.LEFT, padx=16)
        self.lbl_contador_nav = tk.Label(header, text="",
                 font=self._fuentes["small"], bg=BG_CARD, fg=TEXT_MUTED)
        self.lbl_contador_nav.pack(side=tk.LEFT, padx=8)
        tk.Frame(self, bg=BORDER, height=1).pack(fill=tk.X)

        # Info del reporte
        self.frame_info = tk.Frame(self, bg=BG_ITEM, pady=8)
        self.frame_info.pack(fill=tk.X, padx=16, pady=(10, 0))
        self.lbl_audio_info = tk.Label(self.frame_info, text="",
                 font=self._fuentes["small"], bg=BG_ITEM, fg=TEXT_MUTED)
        self.lbl_audio_info.pack(side=tk.LEFT, padx=12)
        self.lbl_paciente = tk.Label(self.frame_info, text="",
                 font=self._fuentes["label"], bg=BG_ITEM, fg=TEXT_ACCENT)
        self.lbl_paciente.pack(side=tk.LEFT, padx=12)
        self.lbl_tipo = tk.Label(self.frame_info, text="",
                 font=self._fuentes["small"], bg=BG_ITEM, fg=TEXT_MUTED)
        self.lbl_tipo.pack(side=tk.LEFT, padx=12)
        self.lbl_guardado = tk.Label(self.frame_info, text="",
                 font=self._fuentes["small"], bg=BG_ITEM, fg=ACCENT_YELLOW)
        self.lbl_guardado.pack(side=tk.RIGHT, padx=12)

        # Barra de audio
        self._construir_audio_bar()

        # Editor
        frame_editor = tk.Frame(self, bg=BG_DARK)
        frame_editor.pack(fill=tk.BOTH, expand=True, padx=16, pady=(8, 0))
        tk.Label(frame_editor,
                 text="Informe — edita el texto si es necesario antes de guardar:",
                 font=self._fuentes["small"], bg=BG_DARK, fg=TEXT_MUTED
                 ).pack(anchor="w", pady=(0, 4))
        frame_texto = tk.Frame(frame_editor, bg=BORDER, bd=1)
        frame_texto.pack(fill=tk.BOTH, expand=True)
        scroll_y = tk.Scrollbar(frame_texto)
        scroll_y.pack(side=tk.RIGHT, fill=tk.Y)
        self.editor = tk.Text(
            frame_texto,
            font=self._fuentes["editor"],
            bg="#0d1117", fg=TEXT_PRIMARY,
            insertbackground="white",
            selectbackground=ACCENT_BLUE,
            wrap=tk.WORD, padx=16, pady=12,
            undo=True, relief="flat",
            yscrollcommand=scroll_y.set
        )
        self.editor.pack(fill=tk.BOTH, expand=True)
        scroll_y.config(command=self.editor.yview)

        # Footer
        tk.Frame(self, bg=BORDER, height=1).pack(fill=tk.X, padx=16, pady=(8, 0))
        footer = tk.Frame(self, bg=BG_DARK, pady=10)
        footer.pack(fill=tk.X, padx=16)

        nav = tk.Frame(footer, bg=BG_DARK)
        nav.pack(side=tk.LEFT)
        self.btn_anterior = self._boton(nav, "◀  Anterior", BG_ITEM, TEXT_MUTED, self._anterior)
        self.btn_anterior.pack(side=tk.LEFT, padx=(0, 6))
        self.btn_siguiente = self._boton(nav, "Siguiente  ▶", BG_ITEM, TEXT_MUTED, self._siguiente)
        self.btn_siguiente.pack(side=tk.LEFT)

        acc = tk.Frame(footer, bg=BG_DARK)
        acc.pack(side=tk.RIGHT)
        self._boton(acc, "💾  Guardar Reporte",
                    ACCENT_GREEN, "white", self._guardar).pack(side=tk.RIGHT, padx=(8, 0))
        self._boton(acc, "📂  Carpeta",
                    BG_ITEM, TEXT_MUTED, self._cambiar_carpeta).pack(side=tk.RIGHT, padx=(8, 0))

        self.lbl_carpeta = tk.Label(footer, text=f"📁 {self._carpeta}",
                 font=self._fuentes["small"], bg=BG_DARK, fg=TEXT_MUTED, cursor="hand2")
        self.lbl_carpeta.pack(side=tk.LEFT, padx=(12, 0))
        self.lbl_carpeta.bind("<Button-1>", lambda e: self._abrir_carpeta())

    def _construir_audio_bar(self):
        """Barra de audio con controles y barra de progreso seekable."""
        bg_audio = "#0d1f3c"

        # Fila 1: controles
        fila1 = tk.Frame(self, bg=bg_audio, pady=6)
        fila1.pack(fill=tk.X, padx=16, pady=(6, 0))

        tk.Label(fila1, text="🎵",
                 font=self._fuentes["body"], bg=bg_audio, fg=TEXT_MUTED
                 ).pack(side=tk.LEFT, padx=(10, 4))

        self.lbl_audio_nombre = tk.Label(fila1, text="—",
                 font=self._fuentes["small"], bg=bg_audio, fg=TEXT_ACCENT)
        self.lbl_audio_nombre.pack(side=tk.LEFT, padx=(0, 12))

        self.btn_play = self._boton_audio(fila1, "▶  Reproducir", ACCENT_BLUE, self._toggle_audio)
        self.btn_play.pack(side=tk.LEFT, padx=(0, 6))

        self.lbl_tiempo = tk.Label(fila1, text="0:00 / 0:00",
                 font=self._fuentes["mono"], bg=bg_audio, fg=TEXT_MUTED)
        self.lbl_tiempo.pack(side=tk.LEFT, padx=8)

        self.lbl_audio_estado = tk.Label(fila1, text="",
                 font=self._fuentes["small"], bg=bg_audio, fg=TEXT_MUTED)
        self.lbl_audio_estado.pack(side=tk.LEFT, padx=6)

        if not self._tiene_ffplay:
            tk.Label(fila1,
                     text="⚠ instala ffmpeg: sudo apt install ffmpeg",
                     font=self._fuentes["small"], bg=bg_audio, fg=ACCENT_YELLOW
                     ).pack(side=tk.RIGHT, padx=10)

        # Fila 2: barra de progreso seekable
        fila2 = tk.Frame(self, bg=bg_audio, pady=2)
        fila2.pack(fill=tk.X, padx=16)

        # Estilo barra de audio
        style = ttk.Style()
        style.configure("Audio.Horizontal.TScale",
                         background=bg_audio,
                         troughcolor="#1a3a6b",
                         sliderlength=14)

        self.var_progreso = tk.DoubleVar(value=0)
        self.scale_audio = ttk.Scale(
            fila2,
            orient=tk.HORIZONTAL,
            variable=self.var_progreso,
            from_=0, to=100,
            style="Audio.Horizontal.TScale",
            command=self._on_scale_move
        )
        self.scale_audio.pack(fill=tk.X, padx=8, pady=(0, 4))
        self.scale_audio.bind("<ButtonPress-1>",   self._on_seek_start)
        self.scale_audio.bind("<ButtonRelease-1>", self._on_seek_end)

        # Separador
        tk.Frame(self, bg=BORDER, height=1).pack(fill=tk.X, padx=16)

    def _boton_audio(self, parent, texto, bg, comando):
        hover = {ACCENT_BLUE: "#4a9eff", ACCENT_RED: "#ff6b6b",
                 ACCENT_YELLOW: "#f0b429", BG_ITEM: BG_HOVER}.get(bg, bg)
        b = tk.Label(parent, text=texto,
                     font=self._fuentes["small"],
                     bg=bg, fg="white",
                     padx=10, pady=4, cursor="hand2")
        b.bind("<Button-1>", lambda e: comando())
        b.bind("<Enter>",    lambda e: b.configure(bg=hover))
        b.bind("<Leave>",    lambda e: b.configure(bg=bg))
        return b

    def _boton(self, parent, texto, bg, fg, comando):
        b = tk.Label(parent, text=texto, font=self._fuentes["label"],
                     bg=bg, fg=fg, padx=14, pady=7, cursor="hand2")
        b.bind("<Button-1>", lambda e: comando())
        b.bind("<Enter>",    lambda e: b.configure(bg=self._hover(bg)))
        b.bind("<Leave>",    lambda e: b.configure(bg=bg))
        return b

    def _hover(self, color):
        return {ACCENT_GREEN: "#5dd879", ACCENT_BLUE: "#4a9eff",
                ACCENT_RED: "#ff6b6b", BG_ITEM: BG_HOVER}.get(color, color)

    # ─── Audio ────────────────────────────────────────────────────────────────
    def _toggle_audio(self):
        if not self._tiene_ffplay:
            messagebox.showwarning("Audio",
                "ffplay no encontrado.\nInstala: sudo apt install ffmpeg", parent=self)
            return
        if self._audio_playing:
            self._detener_audio()
        else:
            self._reproducir_audio()

    def _reproducir_audio(self, desde_seg: float = 0.0):
        rep  = self._reportes[self._indice]
        path = rep.get("path", "")
        if not path or not os.path.isfile(path):
            messagebox.showwarning("Audio", "No se encontró el archivo de audio.", parent=self)
            return
        try:
            self._detener_audio(silencioso=True)

            # Obtener duración si no la tenemos
            if self._audio_duracion == 0.0:
                self._audio_duracion = _duracion_audio(path)
                if self._audio_duracion > 0:
                    self.scale_audio.configure(to=self._audio_duracion)
                    total = _seg_a_tiempo(self._audio_duracion)
                    self.lbl_tiempo.configure(text=f"0:00 / {total}")

            # Lanzar ffplay con -ss para seek
            cmd = ["ffplay", "-nodisp", "-autoexit", "-loglevel", "quiet"]
            if desde_seg > 0:
                cmd += ["-ss", str(desde_seg)]
            cmd.append(path)

            self._ffplay_proc   = subprocess.Popen(cmd,
                stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            self._audio_playing = True
            self._audio_inicio  = time.time() - desde_seg

            self.btn_play.configure(text="⏹  Detener", bg=ACCENT_RED, fg="white")
            self.lbl_audio_estado.configure(text="▶ Reproduciendo", fg=ACCENT_GREEN)
            self._actualizar_progreso()

        except Exception as e:
            messagebox.showerror("Error de audio", str(e), parent=self)

    def _detener_audio(self, silencioso=False):
        if self._ffplay_proc:
            try:
                self._ffplay_proc.terminate()
            except Exception:
                pass
            self._ffplay_proc = None
        self._audio_playing = False
        if hasattr(self, "btn_play") and not silencioso:
            self.btn_play.configure(text="▶  Reproducir", bg=ACCENT_BLUE, fg="white")
            self.lbl_audio_estado.configure(text="", fg=TEXT_MUTED)

    def _actualizar_progreso(self):
        """Actualiza barra y tiempo cada 500ms mientras reproduce."""
        if not self._audio_playing or not self._ffplay_proc:
            return

        # Verificar si ffplay terminó
        if self._ffplay_proc.poll() is not None:
            self._ffplay_proc   = None
            self._audio_playing = False
            self.btn_play.configure(text="▶  Reproducir", bg=ACCENT_BLUE, fg="white")
            self.lbl_audio_estado.configure(text="✅ Terminado", fg=TEXT_MUTED)
            self.var_progreso.set(0)
            total = _seg_a_tiempo(self._audio_duracion)
            self.lbl_tiempo.configure(text=f"0:00 / {total}")
            self.after(2000, lambda: self.lbl_audio_estado.configure(text=""))
            return

        # Calcular posición actual
        if not self._seek_activo:
            pos = time.time() - self._audio_inicio
            pos = min(pos, self._audio_duracion) if self._audio_duracion > 0 else pos
            self.var_progreso.set(pos)
            total = _seg_a_tiempo(self._audio_duracion) if self._audio_duracion > 0 else "--:--"
            self.lbl_tiempo.configure(text=f"{_seg_a_tiempo(pos)} / {total}")

        self.after(500, self._actualizar_progreso)

    # ── Seek (arrastrar la barra) ──────────────────────────────────────────────
    def _on_seek_start(self, event):
        self._seek_activo = True

    def _on_scale_move(self, valor):
        if self._seek_activo and self._audio_duracion > 0:
            pos = float(valor)
            self.lbl_tiempo.configure(
                text=f"{_seg_a_tiempo(pos)} / {_seg_a_tiempo(self._audio_duracion)}")

    def _on_seek_end(self, event):
        self._seek_activo = False
        if self._audio_duracion > 0:
            nueva_pos = self.var_progreso.get()
            if self._audio_playing:
                # Reiniciar ffplay desde la nueva posición
                self._reproducir_audio(desde_seg=nueva_pos)
            else:
                # Solo mover el indicador de tiempo
                total = _seg_a_tiempo(self._audio_duracion)
                self.lbl_tiempo.configure(
                    text=f"{_seg_a_tiempo(nueva_pos)} / {total}")

    # ─── Carga de reporte ─────────────────────────────────────────────────────
    def _cargar_reporte(self, indice: int):
        if not self._reportes:
            return
        self._detener_audio()
        self._audio_duracion = 0.0
        self.var_progreso.set(0)
        if hasattr(self, "lbl_tiempo"):
            self.lbl_tiempo.configure(text="0:00 / 0:00")

        self._indice = indice
        rep = self._reportes[indice]
        res = rep.get("resultado", {}) or {}

        total = len(self._reportes)
        self.lbl_contador_nav.configure(text=f"Reporte {indice + 1} de {total}")

        nombre_audio = rep.get("nombre_audio", "")
        self.lbl_audio_info.configure(text=f"🎵 {nombre_audio}")
        self.lbl_audio_nombre.configure(text=nombre_audio)

        paciente = res.get("nombre_paciente", "") or "Paciente desconocido"
        edad     = res.get("edad", "")
        self.lbl_paciente.configure(
            text=f"👤 {paciente}{f'  ({edad})' if edad else ''}")
        self.lbl_tipo.configure(
            text=f"📋 {res.get('tipo_estudio', 'Tipo no detectado')}")
        self.lbl_guardado.configure(
            text="✅ Guardado" if rep.get("guardado") else "⚠ No guardado",
            fg=ACCENT_GREEN if rep.get("guardado") else ACCENT_YELLOW)

        informe_ia = res.get("informe", rep.get("transcripcion", ""))
        rep["_informe_ia_original"] = informe_ia

        self.editor.configure(state=tk.NORMAL)
        self.editor.delete(1.0, tk.END)
        self.editor.insert(tk.END, informe_ia)

        self.btn_anterior.configure(
            fg=TEXT_PRIMARY if indice > 0 else TEXT_MUTED,
            cursor="hand2" if indice > 0 else "arrow")
        self.btn_siguiente.configure(
            fg=TEXT_PRIMARY if indice < total - 1 else TEXT_MUTED,
            cursor="hand2" if indice < total - 1 else "arrow")

        # Precargar duración en hilo para no bloquear UI
        path = rep.get("path", "")
        if path and os.path.isfile(path):
            def _precargar():
                dur = _duracion_audio(path)
                if dur > 0:
                    self.after(0, lambda: self._set_duracion(dur))
            threading.Thread(target=_precargar, daemon=True).start()

    def _set_duracion(self, dur: float):
        self._audio_duracion = dur
        self.scale_audio.configure(to=dur)
        self.lbl_tiempo.configure(text=f"0:00 / {_seg_a_tiempo(dur)}")

    # ─── Navegación ───────────────────────────────────────────────────────────
    def _anterior(self):
        if self._indice > 0:
            self._sincronizar_editor()
            self._cargar_reporte(self._indice - 1)

    def _siguiente(self):
        if self._indice < len(self._reportes) - 1:
            self._sincronizar_editor()
            self._cargar_reporte(self._indice + 1)

    def _sincronizar_editor(self):
        texto = self.editor.get(1.0, tk.END).strip()
        rep   = self._reportes[self._indice]
        if rep.get("resultado"):
            rep["resultado"]["informe"] = texto
        else:
            rep["resultado"] = {"informe": texto}

    # ─── Guardar ──────────────────────────────────────────────────────────────
    def _guardar(self):
        self._sincronizar_editor()
        rep = self._reportes[self._indice]
        res = rep.get("resultado", {}) or {}

        informe_final = self.editor.get(1.0, tk.END).strip()
        if not informe_final:
            messagebox.showwarning("Aviso", "El informe está vacío.", parent=self)
            return

        try:
            informe_id = InformesRepo.guardar(
                transcripcion_id=rep.get("transcripcion_id"),
                paciente_id=rep.get("paciente_id"),
                contenido=informe_final,
                plantilla_id=None
            )

            # Detectar cambios y guardar ejemplo para IA
            transcripcion_original = rep.get("transcripcion", "")
            informe_ia_original    = rep.get("_informe_ia_original", "")
            hubo_cambios = _textos_diferentes(informe_ia_original, informe_final)

            if transcripcion_original and informe_id:
                calidad = 8 if hubo_cambios else 5
                try:
                    cur = Conexion.cursor()
                    cur.execute(
                        "INSERT INTO ejemplos_ia "
                        "(informe_id, transcripcion, informe_final, calidad) "
                        "VALUES (%s, %s, %s, %s)",
                        (informe_id, transcripcion_original, informe_final, calidad)
                    )
                    cur.close()
                    print(f"[IA] ejemplo guardado — calidad={calidad} — {'con correcciones' if hubo_cambios else 'sin cambios'}")
                except Exception as e:
                    print(f"[ejemplos_ia] Error: {e}")

            nombre_arch   = _nombre_archivo(res)
            ruta_guardada = _guardar_docx(informe_final, nombre_arch, self._carpeta)

            rep["guardado"]   = True
            rep["ruta_docx"]  = ruta_guardada
            rep["informe_id"] = informe_id
            self.lbl_guardado.configure(text="✅ Guardado", fg=ACCENT_GREEN)

            msg = f"Reporte guardado.\n\n📁 {ruta_guardada}"
            msg += "\n\n🧠 Correcciones guardadas — la IA aprenderá de este ejemplo." if hubo_cambios else \
                   "\n\n✅ Informe aceptado sin cambios."
            messagebox.showinfo("Guardado", msg, parent=self)

        except Exception as e:
            messagebox.showerror("Error al guardar", f"No se pudo guardar:\n{e}", parent=self)

    # ─── Carpeta ──────────────────────────────────────────────────────────────
    def _cambiar_carpeta(self):
        nueva = filedialog.askdirectory(
            title="Seleccionar carpeta", initialdir=self._carpeta, parent=self)
        if nueva:
            self._carpeta = nueva
            self.lbl_carpeta.configure(text=f"📁 {self._carpeta}")

    def _abrir_carpeta(self):
        try:
            if sys.platform == "win32":   os.startfile(self._carpeta)
            elif sys.platform == "darwin": subprocess.Popen(["open", self._carpeta])
            else:                          subprocess.Popen(["xdg-open", self._carpeta])
        except Exception:
            pass

    # ─── Cierre ───────────────────────────────────────────────────────────────
    def _cerrar(self):
        self._detener_audio()
        self.grab_release()
        self.destroy()